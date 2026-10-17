"""Build docs/System_Explainer.docx from its editable Markdown source.
Requires Python package python-docx (install in a virtual environment).
Usage: python scripts/build_system_explainer.py
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/System_Explainer.md'
OUTPUT = SOURCE.with_suffix('.docx')

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('153C58')
    code = doc.styles.add_style('Code Block', 1)
    code.font.name = 'Consolas'
    code.font.size = Pt(8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1
    code.paragraph_format.left_indent = Inches(.12)
    sec.header.paragraphs[0].text = 'PACFAS  /  SYSTEM EXPLANATION GUIDE'
    sec.header.paragraphs[0].style = doc.styles['Caption']
    footer = sec.footer.paragraphs[0]
    footer.text = 'Code-based walkthrough • 17 October 2026                                      '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    in_code = False
    first_h1 = True
    for line in SOURCE.read_text().splitlines():
        if line.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(line or ' ', 'Code Block')
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'F0F4F7')
            p._p.get_or_add_pPr().append(shade)
        elif line.startswith('# '):
            if first_h1:
                doc.add_paragraph(line[2:], 'Title')
                first_h1 = False
            else:
                p = doc.add_heading(line[2:], 1)
                p.paragraph_format.page_break_before = True
        elif line.startswith('## '):
            doc.add_heading(line[3:], 2)
        elif line.startswith('• '):
            doc.add_paragraph(line[2:], 'List Bullet')
        elif line.strip():
            doc.add_paragraph(line)
    doc.core_properties.title = 'PACFAS: Explaining the Complicated Parts'
    doc.core_properties.subject = 'Code walkthrough and presentation guide'
    doc.core_properties.keywords = 'OCR, versioning, JWT, transactions, PACFAS'
    doc.save(OUTPUT)
    print(OUTPUT.relative_to(ROOT))

if __name__ == '__main__':
    build()
