"""Build the illustrated Word manual from docs/User_Manual.md.
Requires python-docx and Pillow. Screenshots come from tests/browser.e2e.cjs.
Usage: python scripts/build_user_manual.py
"""
from pathlib import Path
import re
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/User_Manual.md'
OUTPUT = SOURCE.with_suffix('.docx')


def clean(text):
    return text.replace('`', '').replace('**', '')


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('193A52')
    doc.styles['Title'].font.size = Pt(38)
    doc.styles['Heading 1'].font.size = Pt(22)
    doc.styles['Heading 2'].font.size = Pt(14)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.color.rgb = RGBColor.from_string('536170')
    section.header.paragraphs[0].text = 'PACFAS   /   ILLUSTRATED USER MANUAL'
    section.header.paragraphs[0].style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.text = 'Workflow guide • 17 October 2026                                      Page '
    for instr in ['PAGE', 'NUMPAGES']:
        if instr == 'NUMPAGES':
            footer.add_run(' of ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instr)
        footer._p.append(field)
    first_title = True
    lines = SOURCE.read_text().splitlines()
    i = 0
    image_count = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [clean(c.strip()) for c in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, rows[0]):
                cell.text = text
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[1:]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
            for row in table.rows:
                no_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(no_split)
            doc.add_paragraph()
            continue
        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if image:
            caption, filename = image.groups()
            filename = SOURCE.parent / filename
            with Image.open(filename) as im:
                width, height = im.size
            scale = min(6.7 / width, 5.6 / height)
            p = doc.add_paragraph()
            p.alignment = 1
            p.paragraph_format.keep_with_next = True
            inline = p.add_run().add_picture(str(filename), width=Inches(width * scale), height=Inches(height * scale))
            inline._inline.docPr.set('descr', caption)
            cp = doc.add_paragraph(caption, 'Caption')
            cp.alignment = 1
            image_count += 1
        elif line.startswith('# '):
            if first_title:
                doc.add_paragraph(line[2:], 'Title')
                first_title = False
            else:
                h = doc.add_heading(line[2:], 1)
                h.paragraph_format.page_break_before = True
        elif line.startswith('## '):
            doc.add_heading(line[3:], 2)
        elif line.startswith('- '):
            doc.add_paragraph(clean(line[2:]), 'List Bullet')
        elif re.match(r'^\d+\. ', line):
            # Preserve explicit numbering so each task starts at step 1.
            p = doc.add_paragraph(clean(line))
            p.paragraph_format.left_indent = Inches(.12)
        elif line.strip():
            doc.add_paragraph(clean(line))
        i += 1
    assert image_count == 14, f'Expected 14 verified screenshot figures, got {image_count}'
    doc.core_properties.title = 'PACFAS Illustrated User Manual'
    doc.core_properties.subject = 'Browser-tested office workflows with synthetic screenshots'
    doc.core_properties.keywords = 'PACFAS, user manual, documents, cases, OCR'
    doc.save(OUTPUT)
    print(f'Created {OUTPUT.relative_to(ROOT)} with {image_count} screenshots')


if __name__ == '__main__':
    build()
